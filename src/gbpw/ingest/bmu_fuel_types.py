"""
Rebuilds the full BM unit -> fuel type mapping used to identify wind (and
potentially other fuel-type-scoped) BM units -- e.g. for wind curtailment
(see ingest/wind_curtailment.py). Merges two sources:

- The user's manually-downloaded NESO BM Unit Fuel Type spreadsheet ('REG
  FUEL TYPE' column) -- the only source with a real WIND/BATTERY/etc.
  category for most units; not fetchable automatically (NESO's own CMS
  blocks programmatic downloads, confirmed in an earlier round of this
  project), so this ingest takes a local file path rather than a URL.
- The live Elexon reference API (/reference/bmunits/all) -- fills in any
  unit the spreadsheet doesn't cover, and is the only source for
  elexon_bm_unit / lead_party_name / bm_unit_type / generation_capacity_mw
  regardless of fuel type.

Precedence: the spreadsheet's REG FUEL TYPE wins when present -- confirmed
live in an earlier round of this project that the live API's own fuelType
field is ~81% null, and where the two disagree the spreadsheet's dedicated
category is usually the more current/useful one (it's the only source
with a real BATTERY category at all, for instance). The live API's
fuelType only fills in units the spreadsheet has no row for.
"""

from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from ..storage import BmUnitReferenceRow
from . import elexon_bm

# Column headers exactly as they appear in NESO's published spreadsheet --
# confirmed against a real downloaded copy, not guessed (there's no API
# for this file to introspect otherwise).
COL_NESO_BMU_ID = "NESO BMU ID"
COL_REG_FUEL_TYPE = "REG FUEL TYPE"


def _to_float(v: object) -> float | None:
    if v in (None, ""):
        return None
    return float(v)  # type: ignore[arg-type]


def _read_spreadsheet(path: Path) -> dict[str, str]:
    """national_grid_bm_unit -> REG FUEL TYPE, straight from the
    spreadsheet. Blank fuel-type cells are skipped, not stored as "".
    """
    wb = load_workbook(path, read_only=True, data_only=True)
    sheet = wb.worksheets[0]
    rows = sheet.iter_rows(values_only=True)
    header = [str(h).strip() if h is not None else "" for h in next(rows)]
    try:
        id_idx = header.index(COL_NESO_BMU_ID)
        fuel_idx = header.index(COL_REG_FUEL_TYPE)
    except ValueError as e:
        raise ValueError(f"expected columns {COL_NESO_BMU_ID!r}/{COL_REG_FUEL_TYPE!r} not found in {path}") from e

    out: dict[str, str] = {}
    for row in rows:
        bmu_id, fuel_type = row[id_idx], row[fuel_idx]
        if bmu_id and fuel_type:
            out[str(bmu_id).strip()] = str(fuel_type).strip()
    return out


def build_fuel_type_rows(spreadsheet_path: Path) -> list[BmUnitReferenceRow]:
    """Every BM unit known to either source, with a merged fuel_type
    (spreadsheet preferred, live API fills gaps) plus the live API's other
    reference fields. Meant to be passed to
    storage.upsert_bm_unit_reference(..., overwrite_fuel_type=True) --
    this function's whole point is to *recreate* the mapping, not
    incrementally patch it the way the regular reference refresh does.
    """
    spreadsheet_fuel_types = _read_spreadsheet(spreadsheet_path)
    live_records = elexon_bm.fetch_bmu_reference()

    rows: list[BmUnitReferenceRow] = []
    seen: set[str] = set()
    for r in live_records:
        bmu_id = r.get("nationalGridBmUnit")
        if not bmu_id:
            continue
        seen.add(bmu_id)
        rows.append(BmUnitReferenceRow(
            national_grid_bm_unit=bmu_id,
            elexon_bm_unit=r.get("elexonBmUnit"),
            lead_party_name=r.get("leadPartyName"),
            bm_unit_type=r.get("bmUnitType"),
            generation_capacity_mw=_to_float(r.get("generationCapacity")),
            fuel_type=spreadsheet_fuel_types.get(bmu_id) or r.get("fuelType"),
        ))

    # Units the spreadsheet knows about but the live API currently
    # doesn't (e.g. retired/renamed since the spreadsheet was compiled) --
    # confirmed in an earlier round of this project's research to be a
    # real, non-trivial gap (189 units back then), not a hypothetical
    # edge case. Stored with only what the spreadsheet has; the other
    # reference fields stay None until/unless the live API reports that
    # unit under this same ID.
    for bmu_id, fuel_type in spreadsheet_fuel_types.items():
        if bmu_id not in seen:
            rows.append(BmUnitReferenceRow(
                national_grid_bm_unit=bmu_id,
                elexon_bm_unit=None,
                lead_party_name=None,
                bm_unit_type=None,
                generation_capacity_mw=None,
                fuel_type=fuel_type,
            ))
    return rows
